# flask_app.py (à mettre sur PythonAnywhere)

import base64
from flask import Flask, request, jsonify
from flask_cors import CORS
import openpyxl
from io import BytesIO

app = Flask(__name__)

# Origines autorisées : serveurs frontend locaux et GitHub Pages.
CORS(app, origins=[
    'http://127.0.0.1:8000',
    'http://localhost:8000',
    'https://mrchabou.github.io',
])

@app.route('/update-planning', methods=['POST'])
def update_planning():
    try:
        data = request.json
        
        # 1. Décoder le fichier Excel reçu en base64
        excel_data = base64.b64decode(data['file'])
        excel_file = BytesIO(excel_data)
        
        # 2. Charger le classeur avec openpyxl en préservant les styles
        wb = openpyxl.load_workbook(excel_file)
        ws = wb[data['sheet']]
        
        # 3. Récupérer les informations de la requête
        col_idx_target = data['columnIndex'] + 1  # openpyxl est base 1, JS est base 0
        presences = data.get('presences', [])
        
        # Créer un set des noms complets des présents pour une recherche rapide O(1)
        present_players = {f"{p['prenom'].strip().lower()}_{p['nom'].strip().lower()}" for p in presences}
        
        # 4. Parcourir les lignes du fichier Excel pour marquer les présences
        # La logique commence à la ligne 4 (index 3 en JS)
        # On suppose que le nom est en colonne B (2) et le prénom en colonne C (3)
        for row in range(4, ws.max_row + 1):
            stop_text = ' '.join(str(ws.cell(row=row, column=col).value or '') for col in range(1, 5))
            stop_text = stop_text.lower().replace('\u2018', "'").replace('\u2019', "'").replace('\u02bc', "'")
            stop_text = ' '.join(stop_text.split())
            if "liste d'attente" in stop_text:
                break
            
            # Récupérer le nom et prénom de la ligne actuelle
            nom_cell = ws.cell(row=row, column=2).value
            prenom_cell = ws.cell(row=row, column=3).value
            
            # S'il n'y a pas de nom/prénom, on passe à la suite
            if not nom_cell or not prenom_cell:
                continue

            # Construire une clé unique pour le joueur de cette ligne
            player_key = f"{str(prenom_cell).strip().lower()}_{str(nom_cell).strip().lower()}"

            # La cellule cible où marquer la présence
            target_cell = ws.cell(row=row, column=col_idx_target)

            # 5. Mettre à jour la cellule
            is_essai = str(target_cell.value or '').strip().upper() in ('ESSAI', 'ESSAI PRESENT', 'ESSAI ABSENT')
            if is_essai:
                target_cell.value = 'ESSAI PRESENT' if player_key in present_players else 'ESSAI ABSENT'
            elif player_key in present_players:
                target_cell.value = 'V'
            else:
                # Important : vider la cellule si la personne n'est pas marquée présente
                # pour corriger d'éventuelles erreurs précédentes.
                target_cell.value = None

        # 6. Sauvegarder le fichier modifié en mémoire
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 7. Renvoyer le fichier en base64 au frontend
        return jsonify({
            'success': True,
            'file': base64.b64encode(output.read()).decode('utf-8'),
            'filename': data.get('filename', 'planning_mis_a_jour.xlsx')
        })
        
    except Exception as e:
        # En cas d'erreur, renvoyer un message clair
        app.logger.error(f"Erreur lors du traitement: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/health', methods=['GET'])
def health_check():
    """
    Un simple point de terminaison pour vérifier que le serveur est en ligne.
    Le frontend l'utilisera pour déterminer s'il peut préserver les styles.
    """
    return jsonify({'status': 'ok'})
