import base64
import hashlib
from io import BytesIO
from datetime import datetime, timezone
import os
import smtplib
import unittest
from unittest.mock import patch
import openpyxl
from auth_support import authenticated_client
from flask_app import app


class PointageMailTests(unittest.TestCase):
    def setUp(self):
        context = authenticated_client()
        self.client = context.__enter__()
        self.addCleanup(context.__exit__, None, None, None)
        env = dict(BAD_POINTAGE_MAIL_MODE='test', BAD_POINTAGE_MAIL_TEST_TO='test@example.invalid',
                   BAD_POINTAGE_MAIL_PRODUCTION_TO='production@example.invalid',
                   BAD_POINTAGE_MAIL_FROM='sender@example.invalid', BAD_POINTAGE_MAIL_HOST='smtp.example.invalid',
                   BAD_POINTAGE_MAIL_USER='synthetic', BAD_POINTAGE_MAIL_PASSWORD='synthetic')
        self.env = patch.dict(os.environ, env)
        self.env.start()
        self.addCleanup(self.env.stop)
        smtp = patch('pointage_mail.smtplib.SMTP')
        self.smtp = smtp.start().return_value
        self.smtp.send_message.return_value = {}
        self.addCleanup(smtp.stop)
        wb = openpyxl.Workbook()
        wb.active.title = 'Test'
        wb.active.append(['', '', '', 46000])
        wb.active['B4'], wb.active['C4'], wb.active['D4'] = 'Fictif', 'Test', 'ESSAI'
        wb.active['A5'], wb.active['D6'] = "LISTE D'ATTENTE", 'conserver'
        stream = BytesIO()
        wb.save(stream)
        source = stream.getvalue()
        self.data = dict(file=base64.b64encode(source).decode(), sheet='Test', columnIndex=3,
                         presences=[dict(id='P003', nom='Fictif', prenom='Test')],
                         filename='maj_test.xlsx', mode='test', note='Une anomalie synthétique',
                         date='2026-09-29', sourceModified='', sourceHash=hashlib.sha256(source).hexdigest(),
                         participants=1, essais=1, attemptId='12345678-1234-1234-1234-123456789012')

    def send(self):
        return self.client.post('/send-pointage', json=self.data)

    def test_attachment_summary_and_deduplication(self):
        response = self.send()
        self.assertEqual(response.json['state'], 'sent')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        message = self.smtp.send_message.call_args.args[0]
        self.assertEqual(message['To'], 'test@example.invalid')
        self.assertEqual(message['Subject'], '[TEST] Pointage du créneau Test — 29-09-2026')
        body = message.get_body().get_content()
        self.assertTrue(body.startswith('Bonjour Julien,\nVoici le pointage suivant :\n'))
        for line in ('Date : 29-09-2026', 'Participants inscrits : 1', 'Présents : 1',
                     'Absents : 0', 'En essai : 1', 'Fichier : maj_test.xlsx',
                     'Note du responsable de créneau pour réconcilier le fichier ci-joint avec le fichier central :'):
            self.assertIn(line, body)
        business, technical = body.split('Informations techniques\n')
        self.assertIn(self.data['note'], business)
        for label in ('Tentative d’envoi :', 'Source modifiée :', 'SHA-256 source :'):
            self.assertNotIn(label, business)
            self.assertIn(label, technical)
        self.assertTrue(technical.rstrip().endswith(self.data['sourceHash']))
        self.assertEqual(next(message.iter_attachments()).get_filename(), 'maj_test.xlsx')
        self.assertIn(self.data['note'], message.get_body().get_content())
        ws = openpyxl.load_workbook(BytesIO(next(message.iter_attachments()).get_payload(decode=True))).active
        self.assertEqual(ws['D4'].value, 'ESSAI PRESENT')
        self.assertEqual(ws['D6'].value, 'conserver')
        self.assertEqual(self.send().json['state'], 'sent')
        self.smtp.send_message.assert_called_once()

    def test_protected_status_and_send(self):
        for route in ('/mail-status', '/send-pointage'):
            response = app.test_client().open(route, method='GET' if route.endswith('status') else 'POST')
            self.assertEqual(response.status_code, 401)
        self.smtp.send_message.assert_not_called()

    def test_refused_before_smtp(self):
        for key, value in [('mode', 'production'), ('note', 'x' * 2001), ('sourceHash', '0' * 64),
                           ('filename', 'bad\nheader.xlsx'), ('columnIndex', -1)]:
            with self.subTest(key=key):
                original = self.data[key]
                self.data[key] = value
                self.assertEqual(self.send().status_code, 400)
                self.data[key] = original
        self.data['presences'][0]['id'] = 'P999'
        self.assertEqual(self.send().status_code, 400)
        self.smtp.send_message.assert_not_called()

    def test_test_mode_cannot_target_configured_production(self):
        with patch.dict(os.environ, BAD_POINTAGE_MAIL_TEST_TO='production@example.invalid'):
            self.assertEqual(self.send().status_code, 400)
        self.smtp.send_message.assert_not_called()

    def test_production_requires_explicit_matching_mode(self):
        with patch.dict(os.environ, BAD_POINTAGE_MAIL_MODE='production'):
            self.assertEqual(self.send().status_code, 400)
            self.data['mode'] = 'production'
            self.assertEqual(self.send().json['state'], 'sent')
        self.assertEqual(self.smtp.send_message.call_args.args[0]['To'], 'production@example.invalid')
        self.assertEqual(self.smtp.send_message.call_args.args[0]['Subject'],
                         'Pointage du créneau Test — 29-09-2026')

    def test_auth_failure_and_uncertain_transport(self):
        self.smtp.login.side_effect = smtplib.SMTPAuthenticationError(535, b'synthetic')
        self.assertEqual(self.send().json['state'], 'not_sent')
        self.data['attemptId'] = '22345678-1234-1234-1234-123456789012'
        self.smtp.login.side_effect = None
        self.smtp.send_message.side_effect = TimeoutError()
        self.assertEqual(self.send().json['state'], 'uncertain')
        self.assertEqual(self.send().json['state'], 'uncertain')
        self.smtp.send_message.assert_called_once()

    def test_local_attempt_time_respects_summer_and_winter(self):
        for month, expected in ((7, '01-07-2026 à 14:30:00 (CEST)'),
                                (1, '01-01-2026 à 13:30:00 (CET)')):
            with self.subTest(month=month):
                self.data['attemptId'] = f'{month}2345678-1234-1234-1234-123456789012'
                instant = datetime(2026, month, 1, 12, 30, tzinfo=timezone.utc)
                with patch('pointage_mail.datetime', wraps=datetime) as clock:
                    clock.now.side_effect = lambda tz: instant.astimezone(tz)
                    self.assertEqual(self.send().json['state'], 'sent')
                body = self.smtp.send_message.call_args.args[0].get_body().get_content()
                self.assertIn('Tentative d’envoi : ' + expected, body)
