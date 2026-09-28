import base64
from io import BytesIO
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill, Font
from flask_app import app
from auth_support import authenticated_client


class ExportPresenceTests(unittest.TestCase):
    def setUp(self):
        context = authenticated_client()
        self.client = context.__enter__()
        self.addCleanup(context.__exit__, None, None, None)
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Créneau'
        for row in (32, 33, 34, 35, 36, 37, 38):
            self.ws.cell(row, 2, 'Nom')  # Homonymes : seul l'ID distingue les lignes.
            self.ws.cell(row, 3, 'Prénom')
            self.ws.cell(row, 5, 'V')
        self.ws['E36'] = 'ESSAI'
        self.ws['E37'] = 'ESSAI PRESENT'
        self.ws['E38'] = 'ESSAI ABSENT'
        self.ws['B34'].fill = PatternFill('solid', fgColor='FF00FF00')
        # Rouge sombre déjà contrasté sur blanc : ce style doit rester intact.
        self.ws['E35'].font = Font(bold=True, color='FF880000')
        self.ws['E35'].number_format = '@'
        self.ws['A39'] = ' LISTE D’ATTENTE '
        self.ws['B39'] = 'Nom'
        self.ws['C39'] = 'Prénom'
        self.ws['E39'] = 'limite'
        self.ws['B40'] = 'Nom'
        self.ws['C40'] = 'Prénom'
        self.ws['E40'] = 'attente'

    def presence(self, participant_id, **extra):
        return dict(id=participant_id, nom='Nom', prenom='Prénom', **extra)

    def export(self, presences):
        source = BytesIO()
        self.wb.save(source)
        return self.client.post('/update-planning', json=dict(
            file=base64.b64encode(source.getvalue()).decode(), sheet='Créneau',
            columnIndex=4, presences=presences, filename='maj_test.xlsx'))

    def result(self, presences):
        response = self.export(presences)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['filename'], 'maj_test.xlsx')
        return openpyxl.load_workbook(BytesIO(base64.b64decode(response.json['file']))).active

    def test_three_consecutive_repeat_styles_and_new_status(self):
        for nouveau in (True, False):
            presences = [self.presence(f'P{r:03}', nouveauCreneau=nouveau) for r in (32, 33, 34)]
            for _ in range(2):
                ws = self.result(presences)
                self.assertEqual([ws.cell(r, 5).value for r in (33, 34, 35)], ['V'] * 3)
                self.assertIsNone(ws['E32'].value)
                self.assertEqual(ws['E35']._style, self.ws['E35']._style)
                self.assertEqual(ws['B34']._style, self.ws['B34']._style)
                self.assertEqual(ws['E39'].value, 'limite')
                self.assertEqual(ws['E40'].value, 'attente')

    def test_essai_all_markers_present_and_absent(self):
        for presences, expected in (([self.presence(f'P{r:03}') for r in (35, 36, 37)], 'ESSAI PRESENT'), ([], 'ESSAI ABSENT')):
            ws = self.result(presences)
            self.assertEqual([ws.cell(r, 5).value for r in (36, 37, 38)], [expected] * 3)

    def test_legacy_and_normalization(self):
        ws = self.result([dict(nom=' nom ', prenom=' PRÉNOM ')])
        self.assertEqual(ws['E35'].value, 'V')
        ws = self.result([dict(id='P034', nom=' nom ', prenom=' PRÉNOM ')])
        self.assertEqual(ws['E35'].value, 'V')
        self.assertIsNone(ws['E34'].value)

    def test_invalid_id_never_writes_or_falls_back(self):
        invalid = [None, 34, '', 'P34', 'P0034', 'p034', 'P034x', 'P034\n', 'P-34', 'P０３４', 'P002', 'P003', 'P038', 'P039', 'P999', 'P1048576']
        for participant_id in invalid:
            with self.subTest(id=participant_id):
                # Inspect the loaded workbook to ensure even earlier valid entries cannot cause writes.
                with patch('flask_app.openpyxl.load_workbook', return_value=self.wb):
                    response = self.export([self.presence('P032'), self.presence(participant_id)])
                self.assertEqual(response.status_code, 400)
                self.assertNotIn('file', response.json)
                self.assertEqual(self.ws['E36'].value, 'ESSAI')
                self.assertEqual(self.ws['E32'].value, 'V')
        wrong = self.presence('P034'); wrong['nom'] = 'Autre'
        self.assertEqual(self.export([wrong]).status_code, 400)

    def test_ids_above_three_digits(self):
        self.ws['A39'] = None
        self.ws.cell(1001, 2, 'Nom'); self.ws.cell(1001, 3, 'Prénom')
        self.assertEqual(self.result([self.presence('P1000')]).cell(1001, 5).value, 'V')


if __name__ == '__main__':
    unittest.main()
