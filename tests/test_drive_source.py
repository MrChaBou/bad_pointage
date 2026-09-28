"""Drive transport is always simulated: no key or external request."""
import hashlib
import json
from io import BytesIO
import unittest
from unittest.mock import Mock, patch
import openpyxl
import requests
from urllib3.exceptions import ReadTimeoutError
from google.auth import exceptions as google_errors
import drive_source as drive
from flask_app import app
from auth_support import authenticated_client


def fixture():
    wb = openpyxl.Workbook()
    wb.active.append(['', '', '', 46292])
    wb.active.cell(4, 2, 'SYNTHETIQUE')
    wb.active.cell(4, 3, 'Alice')
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class Reply:
    def __init__(self, data=None, status=200, content=b''):
        self.data, self.status_code, self.content = data, status, content
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass
    def json(self):
        return self.data
    def iter_content(self, chunk_size):
        yield self.content


class DriveTests(unittest.TestCase):
    def setUp(self):
        self.bytes = fixture()
        self.meta = dict(name='Créneaux été.xlsx', mimeType=drive.MIME, size=str(len(self.bytes)),
                         version='42', modifiedTime='2026-09-28T10:00:00Z', trashed=False,
                         md5Checksum=hashlib.md5(self.bytes).hexdigest(), capabilities={'canDownload': True})
        self.config = {'BAD_POINTAGE_DRIVE_FILE_ID': 'synthetic-file',
                       'BAD_POINTAGE_GOOGLE_CREDENTIALS_FILE': '/unused/synthetic.json'}
        self.transport = Mock()
        self.transport.__enter__ = Mock(return_value=self.transport)
        self.transport.__exit__ = Mock(return_value=False)
        self.auth = patch.object(drive, 'authorized_session', return_value=self.transport).start()
        self.addCleanup(patch.stopall)
        # Any accidental use of real HTTP fails the test.
        patch('requests.sessions.Session.request', side_effect=AssertionError('Real HTTP forbidden')).start()

    def responses(self, before=None, content=None, after=None):
        self.transport.get.side_effect = [Reply(before or self.meta), Reply(content=self.bytes if content is None else content), Reply(after or self.meta)]

    def error(self, code, status):
        with self.assertRaises(drive.DriveError) as result:
            drive.download(self.config)
        self.assertEqual((result.exception.code, result.exception.status), (code, status))

    def test_original_bytes_shared_drive_fields_and_no_revision_id(self):
        self.responses()
        data, info = drive.download(self.config)
        self.assertEqual(data, self.bytes)
        self.assertEqual(info['sha256'], hashlib.sha256(data).hexdigest())
        self.assertEqual(info['driveVersion'], '42')
        self.assertIsNone(info['headRevisionId'])
        self.assertEqual(self.transport.get.call_count, 3)
        for call in self.transport.get.call_args_list:
            self.assertEqual(call.args[0], 'https://www.googleapis.com/drive/v3/files/synthetic-file')
            self.assertEqual(call.kwargs['params']['supportsAllDrives'], 'true')
            self.assertFalse(call.kwargs['allow_redirects'])
            self.assertGreater(call.kwargs['timeout'], 0)
        self.assertEqual(self.transport.get.call_args_list[0].kwargs['params']['fields'], drive.FIELDS)
        self.assertNotIn('revisionId', drive.FIELDS)
        self.assertEqual(self.transport.get.call_args_list[1].kwargs['params']['alt'], 'media')

    def test_google_errors(self):
        for status, code, output in [(401, 'drive_credentials_invalid', 503), (403, 'drive_access_denied', 502),
                                     (404, 'drive_file_unavailable', 502), (429, 'drive_unavailable', 503),
                                     (500, 'drive_unavailable', 503), (302, 'drive_unavailable', 503)]:
            with self.subTest(status=status):
                self.transport.get.side_effect = [Reply({}, status)]
                self.error(code, output)
        self.transport.get.side_effect = [Reply({'error': {'errors': [{'reason': 'userRateLimitExceeded'}]}}, 403)]
        self.error('drive_unavailable', 503)

    def test_transport_and_credentials_errors(self):
        for exception, code, status in [(requests.Timeout(), 'drive_timeout', 504),
                                       (requests.ConnectionError(), 'drive_unavailable', 503),
                                       (google_errors.RefreshError(), 'drive_credentials_invalid', 503),
                                       (google_errors.RefreshError(retryable=True), 'drive_unavailable', 503),
                                       (requests.ConnectionError(ReadTimeoutError(None, None, 'timeout')), 'drive_timeout', 504),
                                       (google_errors.TransportError(), 'drive_unavailable', 503),
                                       (google_errors.TransportError(requests.Timeout()), 'drive_timeout', 504)]:
            self.transport.get.side_effect = exception
            self.error(code, status)

    def test_missing_configuration(self):
        for value in (None, '', '../other', 'https://example.invalid'):
            self.config['BAD_POINTAGE_DRIVE_FILE_ID'] = value
            self.error('drive_not_configured', 503)
        self.auth.assert_not_called()

    def test_missing_credentials_and_scoped_identity(self):
        # Call original helper via saved function captured below, no actual key generated.
        with self.assertRaises(drive.DriveError) as result:
            original_authorized_session('/nonexistent/bad-pointage-test.json')
        self.assertEqual(result.exception.code, 'drive_not_configured')
        with patch.object(drive.Path, 'is_file', return_value=True), \
             patch.object(drive.service_account.Credentials, 'from_service_account_file') as factory, \
             patch.object(drive, 'AuthorizedSession') as session:
            original_authorized_session('/synthetic.json')
            factory.assert_called_once_with('/synthetic.json', scopes=[drive.SCOPE])
            self.assertEqual(session.call_args.kwargs['refresh_timeout'], 5)
            factory.side_effect = ValueError('private-detail')
            with self.assertRaises(drive.DriveError) as invalid:
                original_authorized_session('/synthetic.json')
            self.assertEqual(str(invalid.exception), 'drive_credentials_invalid')

    def test_metadata_rejections(self):
        for change, code, status in [({'trashed': True}, 'drive_file_unavailable', 502),
                                    ({'capabilities': {'canDownload': False}}, 'drive_access_denied', 502),
                                    ({'mimeType': 'application/vnd.google-apps.spreadsheet'}, 'planning_invalid', 422),
                                    ({'size': '0'}, 'planning_invalid', 422),
                                    ({'size': str(drive.MAX_BYTES + 1)}, 'planning_invalid', 422)]:
            self.responses(before={**self.meta, **change})
            self.error(code, status)

    def test_changed_version(self):
        self.responses(after={**self.meta, 'version': '43'})
        self.error('drive_source_changed', 409)

    def test_invalid_bytes_checksum_and_size(self):
        self.responses(content=self.bytes[:-1])
        self.error('planning_invalid', 422)
        self.responses(content=b'x' * len(self.bytes))
        self.error('planning_invalid', 422)
        bad = b'<html>not an Excel file</html>'
        meta = {**self.meta, 'size': str(len(bad)), 'md5Checksum': None}
        self.responses(before=meta, content=bad, after=meta)
        self.error('planning_invalid', 422)
        self.config['BAD_POINTAGE_DRIVE_MAX_BYTES'] = len(self.bytes)
        self.responses(content=self.bytes + b'x')
        self.error('planning_invalid', 422)

    def test_budget_exhausted(self):
        self.responses()
        with patch.object(drive.time, 'monotonic', side_effect=[0, 31]):
            self.error('drive_timeout', 504)

    def test_route_auth_binary_cors_and_server_target(self):
        with patch.object(drive, 'download') as download:
            response = app.test_client().get('/planning-source')
            self.assertEqual(response.status_code, 401)
            download.assert_not_called()
        self.responses()
        with authenticated_client() as client, patch.dict(app.config, self.config):
            response = client.get('/planning-source?fileId=not-authorized', headers={'Origin': 'http://localhost:8000'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, self.bytes)
        self.assertEqual(response.mimetype, drive.MIME)
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertIn('X-Planning-Metadata', response.headers['Access-Control-Expose-Headers'])
        info = json.loads(response.headers['X-Planning-Metadata'])
        self.assertEqual(info['name'], self.meta['name'])
        self.assertEqual(int(response.headers['Content-Length']), len(self.bytes))
        self.assertNotIn('fileId', info)
        self.assertTrue(all(call.args[0].endswith('/synthetic-file') for call in self.transport.get.call_args_list))

    def test_route_error_sanitized_and_no_store(self):
        with authenticated_client() as client, patch.object(drive, 'download', side_effect=drive.DriveError('drive_credentials_invalid', 503)):
            response = client.get('/planning-source')
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.json, {'error': 'drive_credentials_invalid'})
            self.assertEqual(response.headers['Cache-Control'], 'no-store')


original_authorized_session = drive.authorized_session
