import base64
from copy import copy
from io import BytesIO
import os
from pathlib import Path
import unittest
from auth_support import authenticated_client

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, GradientFill, Side
from openpyxl.writer.theme import theme_xml
from flask_app import (app, planning_theme_colors, planning_color_rgb,
                       planning_contrast, ensure_marker_contrast)


def post_export(workbook, sheet, column, presences):
    data = BytesIO()
    workbook.save(data)
    with authenticated_client() as client:
        result = client.post('/update-planning', json={
            'file': base64.b64encode(data.getvalue()).decode(),
            'sheet': sheet, 'columnIndex': column, 'presences': presences})
    if result.status_code != 200:
        raise AssertionError('Export backend refusé')
    return openpyxl.load_workbook(BytesIO(base64.b64decode(result.json['file'])))


def assert_contrast(test, cell, workbook):
    themes = planning_theme_colors(workbook)
    background = planning_color_rgb(cell.fill.fgColor, themes, workbook._colors)
    foreground = planning_color_rgb(cell.font.color, themes, workbook._colors)
    test.assertIsNotNone(background, 'Couleur du fond résolue')
    test.assertIsNotNone(foreground, 'Couleur de police résolue')
    test.assertGreaterEqual(planning_contrast(background, foreground), 4.5)


class ExportContrastTests(unittest.TestCase):
    def setUp(self):
        self.wb = openpyxl.Workbook()
        # Thème synthétique explicite : theme 0 = noir, theme 1 = blanc.
        self.wb.loaded_theme = theme_xml.replace('lastClr="000000"', 'lastClr="FFFFFF"').replace('<a:srgbClr val="FFFFFF"/>', '<a:srgbClr val="000000"/>').encode()
        # Le thème openpyxl utilise sysClr pour lt1 : remplacer sa valeur également.
        self.wb.loaded_theme = self.wb.loaded_theme.replace(b'val="window" lastClr="FFFFFF"', b'val="window" lastClr="000000"')
        self.ws = self.wb.active
        for row in range(32, 41):
            self.ws.cell(row, 2, 'SYNTHETIQUE')
            self.ws.cell(row, 3, 'Test')
            cell = self.ws.cell(row, 7, 'V')
            cell.fill = PatternFill('solid', fgColor=Color(theme=0))
            cell.font = Font(name='Arial', size=13, bold=True, italic=True, underline='single', color=Color(theme=0))
            cell.border = Border(left=Side(style='thin', color='FFFF0000'))
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.number_format = '@'
        self.ws['G33'].font = copy(self.ws['G35'].font)
        font = copy(self.ws['G33'].font); font.color = Color(theme=1)
        self.ws['G33'].font = font
        for row, value in zip((36, 37, 38), ('ESSAI', 'ESSAI PRESENT', 'ESSAI ABSENT')):
            self.ws.cell(row, 7, value)
        self.ws['A39'] = "LISTE D'ATTENTE"

    def export(self, present_trials):
        ids = [32, 33, 34] + ([35, 36, 37] if present_trials else [])
        presences = [dict(id=f'P{r:03}', nom='SYNTHETIQUE', prenom='Test') for r in ids]
        return post_export(self.wb, self.ws.title, 6, presences)

    def test_markers_only_and_all_other_properties_preserved(self):
        for present_trials in (True, False):
            result = self.export(present_trials)
            ws = result.active
            self.assertEqual([ws.cell(r, 7).value for r in (33, 34, 35)], ['V'] * 3)
            self.assertEqual([ws.cell(r, 7).value for r in (36, 37, 38)],
                             ['ESSAI PRESENT' if present_trials else 'ESSAI ABSENT'] * 3)
            self.assertEqual(ws['G33']._style, self.ws['G33']._style)
            self.assertIsNone(ws['G32'].value)
            for address in ('G32', 'G39', 'G40'):
                self.assertEqual(ws[address]._style, self.ws[address]._style)
            self.assertEqual(ws['G40'].value, 'V')
            for row in range(33, 39):
                source, target = self.ws.cell(row, 7), ws.cell(row, 7)
                assert_contrast(self, target, result)
                self.assertEqual(copy(source.fill), copy(target.fill))
                self.assertEqual(copy(source.border), copy(target.border))
                self.assertEqual(copy(source.alignment), copy(target.alignment))
                self.assertEqual(copy(source.protection), copy(target.protection))
                self.assertEqual(source.number_format, target.number_format)
                expected = copy(source.font); expected.color = target.font.color
                self.assertEqual(expected, copy(target.font))
            self.assertEqual(ws['G35'].font.color, Color(theme=1))

    def test_rgb_indexed_tint_and_unknown_colors(self):
        themes = planning_theme_colors(self.wb)
        for color in (Color(rgb='FF111111'), Color(indexed=0), Color(theme=0, tint=0.05)):
            cell = self.ws['G35']
            cell.fill = PatternFill('solid', fgColor=color)
            cell.font = Font(color=color)
            ensure_marker_contrast(cell, themes)
            assert_contrast(self, cell, self.wb)
        for unknown in (Color(theme=99), Color(auto=True)):
            cell.font = Font(color=unknown)
            original = copy(cell._style)
            ensure_marker_contrast(cell, themes)
            self.assertEqual(cell._style, original)
        # Thème Office habituel : ne jamais assimiler automatiquement theme 0 au noir.
        self.wb.loaded_theme = theme_xml.encode()
        cell.fill = PatternFill('solid', fgColor=Color(theme=1))
        cell.font = Font(color=Color(theme=1))
        ensure_marker_contrast(cell, planning_theme_colors(self.wb))
        self.assertEqual(cell.font.color, Color(theme=0))
        assert_contrast(self, cell, self.wb)

    def test_gradient_is_preserved_without_assuming_a_uniform_background(self):
        cell = self.ws['G35']
        cell.fill = GradientFill(stop=('FF000000', 'FFFFFFFF'))
        style = copy(cell._style)
        ensure_marker_contrast(cell, planning_theme_colors(self.wb))
        self.assertEqual(cell._style, style)

@unittest.skipUnless(os.environ.get('BAD_POINTAGE_REFERENCE'), 'Classeur réel optionnel non fourni')
class LocalReferenceContrastTests(unittest.TestCase):
    def test_real_export_without_personal_output(self):
        source = openpyxl.load_workbook(Path(os.environ['BAD_POINTAGE_REFERENCE']))
        sheet = source['Vendredi 18h00 à 20h00']
        presences = [dict(id=f'P{row - 1:03}', nom=str(sheet.cell(row, 2).value).strip(),
                          prenom=str(sheet.cell(row, 3).value).strip()) for row in (33, 34, 35)]
        result = post_export(source, sheet.title, 6, presences)
        for row in (33, 34, 35):
            cell = result[sheet.title].cell(row, 7)
            self.assertEqual(cell.value, 'V', 'Marqueur attendu sur une ligne témoin')
            assert_contrast(self, cell, result)
